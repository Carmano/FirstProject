import docx
import ParserSibGUTI


def group_schedule(file):
    names_group = []
    days_of_the_week = ('monday', 'tuesday', 'wednesday', 'thursday', 'friday', 'saturday')
    doc = docx.Document(file)
    table = doc.tables[1]
    count_column = len(table.columns)
    for i in range(2, count_column):
        print(table.cell(1, i).text)


def main():
    file = 'files/ИИ 051, 052.docx'
    group_schedule(file)


if __name__ == '__main__':
    main()
